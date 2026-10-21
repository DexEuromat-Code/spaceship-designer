"""Builds Spaceship_Designer_Test_Results.xlsx from the BDD markdown test files.

Usage:  python tests/manual/build_results_sheet.py
Each "Then" line with an ID (e.g. **BLK-01**) becomes one row to fill in (PASS / FAIL + comment).
"""
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
OUT = HERE / "Spaceship_Designer_Test_Results.xlsx"

RE_FEATURE = re.compile(r"^# Feature:\s*(.+)$")
RE_GIVEN = re.compile(r"^## (Given .+)$")
RE_WHEN = re.compile(r"^### (When .+)$")
RE_THEN = re.compile(r"^- \[ \] \*\*([A-Z]{3}-\d{2})\*\*\s*—\s*(Then .+)$")


def clean(text: str) -> str:
    """Strip markdown emphasis/code marks so cells read as plain text."""
    return re.sub(r"[*`]", "", text).strip()


def parse_tests():
    rows, features = [], []
    for md in sorted(HERE.glob("[0-9][0-9]-*.md")):
        feature = given = when = None
        for line in md.read_text(encoding="utf-8").splitlines():
            if m := RE_FEATURE.match(line):
                feature = clean(m.group(1))
                features.append((feature, md.name))
            elif m := RE_GIVEN.match(line):
                given, when = clean(m.group(1)), None
            elif m := RE_WHEN.match(line):
                when = clean(m.group(1))
            elif m := RE_THEN.match(line):
                if not (feature and given and when):
                    raise ValueError(f"{md.name}: {m.group(1)} is not under Feature/Given/When")
                rows.append((m.group(1), feature, given, when, clean(m.group(2)), md.name))
    ids = [r[0] for r in rows]
    dupes = {i for i in ids if ids.count(i) > 1}
    if dupes:
        raise ValueError(f"Duplicate test IDs: {sorted(dupes)}")
    return rows, features


FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_MUTED = Font(name=FONT, size=9, italic=True, color="555555")
FILL_HEAD = PatternFill("solid", fgColor="1F2A38")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")   # cells the tester fills in
FILL_BAND = PatternFill("solid", fgColor="F3F5F8")
FILL_PASS = PatternFill("solid", fgColor="C6EFCE")
FILL_FAIL = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="BFC5CD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top")


def build_results(ws, rows):
    ws.title = "Test Results"
    ws["A1"] = "Spaceship Designer — Manual Test Results"
    ws["A1"].font = F_TITLE

    meta = ["Tester", "Test date", "Browser & version", "App build / file date"]
    for i, label in enumerate(meta, start=3):
        ws.cell(row=i, column=1, value=label).font = F_BOLD
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
        c = ws.cell(row=i, column=2)
        c.fill, c.font, c.border = FILL_INPUT, F_BASE, BORDER
        ws.cell(row=i, column=3).border = BORDER

    ws["E3"] = "How to fill in"
    ws["E3"].font = F_BOLD
    legend = [
        "Edit only the yellow cells: the run details above, Result and Tester comment.",
        "Result: choose PASS or FAIL from the drop-down. Leave empty if the test was not run.",
        'Example: Result = FAIL, comment = "Counter stays 0/25 after typing; Edge 129, Win 11. Steps: …"',
        "Test IDs match the markdown files in tests/manual (e.g. BLK-01 → 01-block-creation.md).",
    ]
    for i, text in enumerate(legend, start=4):
        ws.cell(row=i, column=5, value=text).font = F_MUTED

    header_row = 9
    headers = ["ID", "Feature", "Given", "When", "Then (expected result)", "Result", "Tester comment"]
    widths = [10, 22, 38, 40, 60, 11, 45]
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[header_row].height = 22

    first = header_row + 1
    band, prev_feature = False, None
    for r, (tid, feature, given, when, then, _src) in enumerate(rows, start=first):
        if feature != prev_feature:
            band, prev_feature = not band, feature
        values = [tid, feature, given, when, then, None, None]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.font, c.border, c.alignment = F_BASE, BORDER, WRAP_TOP
            if col in (6, 7):
                c.fill = FILL_INPUT
            elif band:
                c.fill = FILL_BAND
        ws.cell(row=r, column=1).font = F_BOLD
        ws.cell(row=r, column=6).alignment = CENTER
    last = first + len(rows) - 1

    dv = DataValidation(type="list", formula1='"PASS,FAIL"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Invalid result",
                        error="Choose PASS or FAIL from the list.")
    ws.add_data_validation(dv)
    dv.add(f"F{first}:F{last}")

    rng = f"F{first}:F{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$F{first}="PASS"'], fill=FILL_PASS,
                                                   font=Font(name=FONT, bold=True, color="006100")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$F{first}="FAIL"'], fill=FILL_FAIL,
                                                   font=Font(name=FONT, bold=True, color="9C0006")))

    ws.freeze_panes = ws.cell(row=first, column=2)
    ws.auto_filter.ref = f"A{header_row}:G{last}"
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return first, last


def build_summary(ws, features, first, last):
    ws["A1"] = "Test Summary"
    ws["A1"].font = F_TITLE
    ws["A2"] = "Calculated automatically from the Test Results sheet."
    ws["A2"].font = F_MUTED

    headers = ["Feature", "Test file", "Tests", "PASS", "FAIL", "Not run", "Pass rate (of run)"]
    widths = [34, 26, 9, 9, 9, 10, 18]
    hr = 4
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=hr, column=col, value=h)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
        ws.column_dimensions[get_column_letter(col)].width = w

    res = "'Test Results'"
    feat_rng = f"{res}!$B${first}:$B${last}"
    result_rng = f"{res}!$F${first}:$F${last}"
    r = hr
    for r, (feature, src) in enumerate(features, start=hr + 1):
        ws.cell(row=r, column=1, value=feature)
        ws.cell(row=r, column=2, value=src)
        ws.cell(row=r, column=3, value=f"=COUNTIF({feat_rng},A{r})")
        ws.cell(row=r, column=4, value=f'=COUNTIFS({feat_rng},A{r},{result_rng},"PASS")')
        ws.cell(row=r, column=5, value=f'=COUNTIFS({feat_rng},A{r},{result_rng},"FAIL")')
        ws.cell(row=r, column=6, value=f"=C{r}-D{r}-E{r}")
        ws.cell(row=r, column=7, value=f'=IF(D{r}+E{r}=0,"-",D{r}/(D{r}+E{r}))')
    end = r
    tr = end + 1
    ws.cell(row=tr, column=1, value="TOTAL")
    for col in range(3, 7):
        L = get_column_letter(col)
        ws.cell(row=tr, column=col, value=f"=SUM({L}{hr + 1}:{L}{end})")
    ws.cell(row=tr, column=7, value=f'=IF(D{tr}+E{tr}=0,"-",D{tr}/(D{tr}+E{tr}))')

    for row in ws.iter_rows(min_row=hr + 1, max_row=tr, max_col=7):
        for c in row:
            c.border = BORDER
            c.font = F_BOLD if c.row == tr else F_BASE
            if c.column >= 3:
                c.alignment = Alignment(horizontal="center")
            if c.column == 7:
                c.number_format = "0%"
    for col in range(1, 8):
        ws.cell(row=tr, column=col).fill = FILL_BAND

    ws.cell(row=tr + 2, column=1, value="Overall status").font = F_BOLD
    ws.cell(row=tr + 2, column=2,
            value=f'=IF(E{tr}>0,"FAILED",IF(F{tr}>0,"IN PROGRESS","PASSED"))').font = F_BOLD
    ws.conditional_formatting.add(f"B{tr + 2}", FormulaRule(formula=[f'B{tr + 2}="FAILED"'], fill=FILL_FAIL))
    ws.conditional_formatting.add(f"B{tr + 2}", FormulaRule(formula=[f'B{tr + 2}="PASSED"'], fill=FILL_PASS))
    ws.cell(row=tr + 3, column=1,
            value="FAILED = at least one FAIL · IN PROGRESS = tests not run yet · PASSED = every test passed").font = F_MUTED


def main():
    rows, features = parse_tests()
    wb = Workbook()
    first, last = build_results(wb.active, rows)
    build_summary(wb.create_sheet("Summary"), features, first, last)
    wb.calculation.fullCalcOnLoad = True   # Excel computes the summary on open
    wb.save(OUT)
    print(f"{len(rows)} tests in {len(features)} features -> {OUT.name}")


if __name__ == "__main__":
    main()
